"""Read Excel cell values without evaluating formulas or executing macros."""
from contextlib import ExitStack
from datetime import date,datetime,time
from pathlib import Path
import re
from openpyxl import load_workbook


class ExcelImportError(ValueError):pass


def cell_text(cell):
    value=cell.value
    if value is None:return ''
    if isinstance(value,bool):return 'TRUE' if value else 'FALSE'
    if isinstance(value,(datetime,date,time)):return value.isoformat()
    if isinstance(value,(int,float)):
        # Common identifier format: preserve leading zeroes in e.g. 0000.
        if re.fullmatch(r'0+',cell.number_format or '') and float(value).is_integer():
            number=int(value);return ('-' if number<0 else '')+str(abs(number)).zfill(len(cell.number_format))
        if isinstance(value,float) and value.is_integer():return str(int(value))
    return str(value)


class ExcelBook:
    def __init__(self,path):
        self.path=Path(path);self._stack=ExitStack()
        try:
            if self.path.suffix.lower()!='.xlsx':
                raise ExcelImportError('Excelの.xlsxファイルを指定してください。.xlsはExcelで.xlsxとして保存してください。')
            self.source=load_workbook(self.path,read_only=True,data_only=False,keep_links=False)
            self._stack.callback(self.source.close)
            self.sheets=[sheet.title for sheet in self.source.worksheets]
            if not self.sheets:raise ExcelImportError('読み込めるワークシートがありません。')
        except Exception as error:
            self.close()
            if isinstance(error,ExcelImportError):raise
            raise ExcelImportError(f'Excelを開けません: {error}') from error

    def __enter__(self):return self
    def __exit__(self,*args):self.close()
    def close(self):self._stack.close()

    def read(self,sheet_name):
        try:
            sheet=self.source[sheet_name]
            rows,columns=sheet.max_row or 1,sheet.max_column or 1
            if rows>10000 or columns>256 or rows*columns>200000:
                raise ExcelImportError('表は10,000行・256列・200,000セル以内にしてください。不要な行・列の書式も削除してください。')
            cached=load_workbook(self.path,read_only=True,data_only=True,keep_links=False)
            try:
                data=[];last_row=last_column=0
                originals=sheet.iter_rows(min_row=1,min_col=1,max_row=rows,max_col=columns)
                values=cached[sheet_name].iter_rows(min_row=1,min_col=1,max_row=rows,max_col=columns)
                for row_number,(source_row,value_row) in enumerate(zip(originals,values),1):
                    result=[]
                    for column,(original,value) in enumerate(zip(source_row,value_row),1):
                        if original.data_type=='f' and value.value is None:
                            raise ExcelImportError(f'{sheet_name}!{original.coordinate}: 数式の計算済み値がありません。Excelで再計算して保存してください。')
                        if value.data_type=='e':
                            raise ExcelImportError(f'{sheet_name}!{value.coordinate}: Excelのセルエラー {value.value} を修正してください。')
                        text=cell_text(value);result.append(text)
                        if text!='':last_row=row_number;last_column=max(last_column,column)
                    data.append(result)
            finally:cached.close()
            if not last_row:raise ExcelImportError('選択したシートに値がありません。')
            # Keep leading/interior empty cells and rows: origin is always A1.
            return [row[:last_column] for row in data[:last_row]]
        except ExcelImportError:raise
        except Exception as error:raise ExcelImportError(f'Excelを読み込めません: {error}') from error
