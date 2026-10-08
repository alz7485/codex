import copy
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree
from datetime import date
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication,QDialog
from e3d_designer.excel_import import ExcelBook,ExcelImportError
from e3d_designer.excel_dialog import ExcelSheetDialog
from e3d_designer.quick_editor import ItemsDialog,MiniProperties
from e3d_designer.model import Form,Gadget
from e3d_designer.mac_import import import_mac
from e3d_designer.app import Window


class ExcelReaderTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.path=Path(self.temp.name)/'values.xlsx'
    def tearDown(self):self.temp.cleanup()
    def save(self,rows):
        book=Workbook();sheet=book.active;sheet.title='値'
        for row in rows:sheet.append(row)
        book.save(self.path);book.close()
    def read(self,sheet='値'):
        with ExcelBook(self.path) as book:return book.read(sheet)

    def test_a1_origin_internal_blanks_and_formatted_trailing_cells(self):
        self.save([['名前','値',None],[' MiXeD ',0,None],[None,None,None],['日本語',False,'last']])
        book=__import__('openpyxl').load_workbook(self.path);book.active['F9'].number_format='0.00';book.save(self.path);book.close()
        self.assertEqual(self.read(),[['名前','値',''],[' MiXeD ','0',''],['','',''],['日本語','FALSE','last']])

    def test_does_not_shift_leading_empty_rows_or_columns(self):
        self.save([[None,None],[None,'B2']]);self.assertEqual(self.read(),[['',''],['','B2']])

    def test_sheet_selection_and_empty_sheet(self):
        book=Workbook();book.active.title='空';sheet=book.create_sheet('Data');sheet.append(['A','B']);book.save(self.path);book.close()
        with ExcelBook(self.path) as reader:
            self.assertEqual(reader.sheets,['空','Data']);self.assertEqual(reader.read('Data'),[['A','B']])
            with self.assertRaisesRegex(ExcelImportError,'値がありません'):reader.read('空')
            with self.assertRaisesRegex(ExcelImportError,'読み込めません'):reader.read('missing')

    def test_types_dates_and_simple_leading_zero_format(self):
        book=Workbook();sheet=book.active;sheet.title='値'
        sheet.append(['0012',12,1.5,True,date(2026,10,8),-7]);sheet['B1'].number_format='0000';sheet['F1'].number_format='000'
        book.save(self.path);book.close()
        self.assertEqual(self.read()[0],['0012','0012','1.5','TRUE','2026-10-08T00:00:00','-007'])

    def test_formula_without_cache_is_not_silently_imported_as_empty(self):
        self.save([['A','=1+2']])
        with self.assertRaisesRegex(ExcelImportError,'B1.*再計算'):self.read()

    def test_cached_formula_value_and_error(self):
        self.save([['A','=1+2']])
        with zipfile.ZipFile(self.path) as archive:files={name:archive.read(name) for name in archive.namelist()}
        key='xl/worksheets/sheet1.xml';root=ElementTree.fromstring(files[key]);ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        root.find(".//s:c[@r='B1']/s:v",ns).text='3';files[key]=ElementTree.tostring(root)
        with zipfile.ZipFile(self.path,'w') as archive:
            for name,value in files.items():archive.writestr(name,value)
        self.assertEqual(self.read(),[['A','3']])
        self.save([['A','#DIV/0!']])
        with self.assertRaisesRegex(ExcelImportError,'B1.*セルエラー'):self.read()

    def test_bad_file_and_legacy_extension_report_actionable_error(self):
        self.path.write_bytes(b'not excel')
        with self.assertRaisesRegex(ExcelImportError,'開けません'):ExcelBook(self.path)
        with self.assertRaisesRegex(ExcelImportError,'.xls.*.xlsx'):ExcelBook(self.path.with_suffix('.xls'))

    def test_oversized_sheet_is_rejected_without_truncation(self):
        book=Workbook();book.active.title='値';book.active['A10001']='last';book.save(self.path);book.close()
        with self.assertRaisesRegex(ExcelImportError,'10,000'):self.read()


class ExcelEditorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.folder=Path(self.temp.name);self.dialogs=[]
    def tearDown(self):
        for dialog in self.dialogs:dialog.reject();dialog.deleteLater()
        self.app.processEvents();self.temp.cleanup()
    def editor(self,**kwargs):
        dialog=ItemsDialog(None,Gadget(name='Choice',**kwargs));self.dialogs.append(dialog);return dialog
    def grid(self,d):return [[d.table.item(r,c).text() if d.table.item(r,c) else '' for c in range(d.table.columnCount())] for r in range(d.table.rowCount())]

    def test_option_commands_preserve_spaces_quotes_and_server_dollar(self):
        d=self.editor(kind='option',items=['old'],item_commands=['old'])
        rows=[["It's a choice",'$M /$Server/File.mac'],[' MiXeD ','!this.val = |a b|']]
        d.apply_excel_data(rows);d.accept()
        self.assertEqual(d.gadget.items,[r[0] for r in rows]);self.assertEqual(d.gadget.item_commands,[r[1] for r in rows]);self.assertEqual(d.gadget.item_values,[])
        form=Form(gadgets=[d.gadget]);form.validate();loaded=import_mac(form.pml()).form
        self.assertEqual(loaded.gadgets[0].items,d.gadget.items);self.assertEqual(loaded.gadgets[0].item_commands,d.gadget.item_commands)

    def test_option_values_combo_and_simple_list(self):
        for kind in ('option','combo','list'):
            d=self.editor(kind=kind,items=['old'],item_values=['old'])
            d.apply_excel_data([['表示名','値'],['A','001'],['B','0']],first_row_header=True);d.accept()
            self.assertEqual(d.gadget.items,['A','B']);self.assertEqual(d.gadget.item_values,['001','0']);self.assertEqual(d.gadget.item_commands,[])
            form=Form(gadgets=[d.gadget]);form.validate();loaded=import_mac(form.pml()).form.gadgets[0]
            self.assertEqual(loaded.items,['A','B']);self.assertEqual(loaded.item_values,['001','0'])

    def test_optional_b_column_clears_previous_values(self):
        d=self.editor(kind='combo',items=['old'],item_values=['old'])
        d.apply_excel_data([['A'],['B']]);d.accept();self.assertEqual(d.gadget.items,['A','B']);self.assertEqual(d.gadget.item_values,[])

    def test_list_table_headings_rows_and_roundtrip(self):
        d=self.editor(kind='list',items=['old']);data=[['名前','値','型'],['A','001','TEXT'],['B','','REAL']]
        d.apply_excel_data(data,table_mode=True,first_row_header=True);d.accept()
        self.assertEqual(d.gadget.list_mode,'TABLE');self.assertEqual(d.gadget.headings,data[0]);self.assertEqual(d.gadget.rows,data[1:])
        form=Form(gadgets=[d.gadget]);form.validate();loaded=import_mac(form.pml()).form.gadgets[0]
        self.assertEqual(loaded.headings,data[0]);self.assertEqual(loaded.rows,data[1:])
        self.assertEqual(Form.loads(form.dumps()).dumps(),form.dumps())

    def test_list_table_without_header_keeps_a1_as_data(self):
        d=self.editor(kind='list');d.apply_excel_data([['A','1','X'],['B','2','Y']],True,False);d.accept()
        self.assertEqual(d.gadget.headings,['列1','列2','列3']);self.assertEqual(d.gadget.rows[0],['A','1','X'])

    def test_invalid_import_keeps_current_unsaved_table_and_mode(self):
        d=self.editor(kind='option',items=['draft'],item_commands=['command']);before=self.grid(d)
        for rows in ([['A','B','C']],[['A\nB','X']],[['NUL\x00','X']],[['$bad','X']]):
            with self.assertRaises(ValueError):d.apply_excel_data(rows)
            self.assertEqual(self.grid(d),before);self.assertEqual(d.mode.currentIndex(),0)

    def test_import_replaces_mode_cache_without_resurrecting_old_data(self):
        d=self.editor(kind='list',items=['OLD'],item_values=['99']);d.mode.setCurrentIndex(1);d.mode.setCurrentIndex(0)
        d.apply_excel_data([['A','B'],['NEW','1']],True,True)
        d.mode.setCurrentIndex(0);self.assertEqual(self.grid(d),[['NEW','1']])
        d.mode.setCurrentIndex(1);self.assertEqual(self.grid(d),[['A','B'],['NEW','1']])

    def test_excel_picker_defaults_and_cancel_preserve_draft(self):
        for is_list in (True,False):
            d=ExcelSheetDialog(None,['First','Second'],is_list);self.dialogs.append(d)
            self.assertEqual(d.table_mode(),is_list);self.assertEqual(d.header.isChecked(),is_list)
        d=self.editor(kind='option',items=['Draft']);before=self.grid(d)
        with patch('e3d_designer.quick_editor.QFileDialog.getOpenFileName',return_value=('','')):d.excel_button.click()
        self.assertEqual(self.grid(d),before)

    def test_actual_import_button_with_sheet_selection(self):
        path=self.folder/'sample.xlsx';book=Workbook();book.active.append(['wrong']);book.create_sheet('Data').append(['Caption','001']);book.save(path);book.close()
        d=self.editor(kind='combo',items=['Draft']);d.show();self.app.processEvents()
        def choose_sheet(dialog):dialog.sheet.setCurrentText('Data');return QDialog.Accepted
        with patch('e3d_designer.quick_editor.QFileDialog.getOpenFileName',return_value=(str(path),'Excel (*.xlsx)')), patch('e3d_designer.excel_dialog.ExcelSheetDialog.exec',choose_sheet):
            QTest.mouseClick(d.excel_button,Qt.LeftButton)
        self.assertEqual(self.grid(d),[['Caption','001']]);self.assertEqual(d.gadget.items,['Draft'])
        d.accept();self.assertEqual(d.gadget.items,['Caption']);self.assertEqual(d.gadget.item_values,['001'])

    def test_failed_file_or_cancelled_settings_leave_editor_unchanged(self):
        d=self.editor(kind='combo',items=['Draft']);before=self.grid(d)
        d.load_excel(str(self.folder/'missing.xlsx'));self.assertTrue(d.error.text());self.assertEqual(self.grid(d),before)
        path=self.folder/'valid.xlsx';book=Workbook();book.active.append(['New']);book.save(path);book.close()
        with patch('e3d_designer.excel_dialog.ExcelSheetDialog.exec',lambda dialog:QDialog.Rejected):d.load_excel(path)
        self.assertEqual(self.grid(d),before)

    def test_outer_editor_cancel_and_application_commit_undo(self):
        original=Form(gadgets=[Gadget(kind='list',name='Rows',list_mode='TABLE',headings=['Old'],rows=[['old']])])
        outer=MiniProperties(None,original,0);self.dialogs.append(outer)
        def fill_items(dialog):
            dialog.apply_excel_data([['Name','Value'],['A','001']],True,True);dialog.accept();return dialog.result()
        with patch('e3d_designer.quick_editor.ItemsDialog.exec',fill_items):outer.edit_items()
        outer.reject();self.assertEqual(original.gadgets[0].headings,['Old'])
        window=Window(settings_path=self.folder/'settings.json');window.form=copy.deepcopy(original);window.refresh();before=window.form.dumps()
        def commit(dialog):
            with patch('e3d_designer.quick_editor.ItemsDialog.exec',fill_items):dialog.edit_items()
            dialog.accept();return dialog.result()
        try:
            with patch('e3d_designer.quick_editor.MiniProperties.exec',commit):window.edit_object_properties('Rows')
            self.assertEqual(window.form.named('Rows').rows,[['A','001']]);self.assertEqual(len(window.history),1)
            after=window.form.dumps();window.undo();self.assertEqual(window.form.dumps(),before);window.redo();self.assertEqual(window.form.dumps(),after)
        finally:window.dirty=False;window.close();self.app.processEvents()

