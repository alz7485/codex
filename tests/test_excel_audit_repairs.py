import copy
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from unittest.mock import patch
from openpyxl import Workbook
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication,QDialog,QTableWidgetItem,QDialogButtonBox,QLineEdit
from e3d_designer.excel_import import ExcelBook,ExcelImportError
from e3d_designer.quick_editor import ItemsDialog,MiniProperties
from e3d_designer.item_editing import validate_item_change
from e3d_designer.model import Form,Gadget
from e3d_designer.mac_import import import_mac
from e3d_designer.clipboard import clone_subtree
from e3d_designer.app import Window


class ExcelReadAuditTests(unittest.TestCase):
    def setUp(self):self.temp=tempfile.TemporaryDirectory();self.path=Path(self.temp.name)/'table.xlsx'
    def tearDown(self):self.temp.cleanup()
    def save(self,rows):
        book=Workbook()
        for row in rows:book.active.append(row)
        book.save(self.path);book.close()
    def xml(self,change):
        with zipfile.ZipFile(self.path) as archive:files={name:archive.read(name) for name in archive.namelist()}
        key='xl/worksheets/sheet1.xml';root=ET.fromstring(files[key]);change(root,{'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'});files[key]=ET.tostring(root)
        with zipfile.ZipFile(self.path,'w') as archive:
            for name,data in files.items():archive.writestr(name,data)
    def read(self):
        with ExcelBook(self.path) as reader:return reader.read('Sheet')

    def test_small_missing_and_excessive_dimensions_do_not_lose_data(self):
        for dimension in ('A1',None,'A1:XFD1048576'):
            self.save([['A','B'],['C','D']])
            def change(root,ns):
                node=root.find('s:dimension',ns)
                if dimension is None:root.remove(node)
                else:node.set('ref',dimension)
            self.xml(change);self.assertEqual(self.read(),[['A','B'],['C','D']])

    def test_unknown_dimensions_preserve_ragged_rows_and_interior_empty_rows(self):
        self.save([['A'],[],['B',None,'C'],['D']]);self.xml(lambda root,ns:root.remove(root.find('s:dimension',ns)))
        self.assertEqual(self.read(),[['A','',''],['','',''],['B','','C'],['D','','']])

    def test_cached_empty_string_formula_is_an_empty_cell(self):
        self.save([['A','=""'],['B',1]])
        self.xml(lambda root,ns:root.find(".//s:c[@r='B1']",ns).set('t','str'))
        self.assertEqual(self.read(),[['A',''],['B','1']])

    def test_empty_cached_formula_retains_column_and_single_empty_cell(self):
        for rows,expected,cell in ([['A','=""']],[['A','']],'B1'),([['=""']],[['']],'A1'):
            self.save(rows);self.xml(lambda root,ns:root.find(f".//s:c[@r='{cell}']",ns).set('t','str'))
            self.assertEqual(self.read(),expected)

    def test_missing_dimensions_still_enforce_actual_limits(self):
        for coordinate in ('A10001','IW1'):
            book=Workbook();book.active[coordinate]='last';book.save(self.path);book.close()
            self.xml(lambda root,ns:root.remove(root.find('s:dimension',ns)))
            with self.assertRaisesRegex(ExcelImportError,'10,000'):self.read()

    def test_repeated_reads_and_close_release_shared_file(self):
        self.save([['A','B']])
        with ExcelBook(self.path) as reader:
            self.assertEqual(reader.read('Sheet'),[['A','B']]);self.assertEqual(reader.read('Sheet'),[['A','B']]);self.assertFalse(reader._file.closed)
        self.assertTrue(reader._file.closed)

    def test_both_formula_and_cached_reads_use_the_same_open_file(self):
        self.save([['Original','1']])
        other=Path(self.temp.name)/'other.xlsx';book=Workbook();book.active.append(['Replacement','2']);book.save(other);book.close()
        with ExcelBook(self.path) as reader:
            reader.path=other
            self.assertEqual(reader.read('Sheet'),[['Original','1']])


class ExcelEditorAuditTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def setUp(self):self.dialogs=[];self.temp=tempfile.TemporaryDirectory();self.folder=Path(self.temp.name)
    def tearDown(self):
        for dialog in self.dialogs:dialog.reject();dialog.deleteLater()
        self.app.clipboard().clear();self.app.processEvents();self.temp.cleanup()
    def editor(self,g,**kwargs):
        d=ItemsDialog(None,g,**kwargs);self.dialogs.append(d);return d
    def window(self,form):
        w=Window(settings_path=self.folder/'settings.json');w.form=form;w.selected=0;w.refresh();return w

    def test_empty_option_real_values_reopen_and_roundtrip_without_becoming_commands(self):
        d=self.editor(Gadget(kind='option',name='Choice',items=['Old'],item_values=['1']))
        d.apply_excel_data([['A',''],['B','']]);d.accept();self.assertEqual(d.gadget.item_values,['','']);self.assertEqual(d.gadget.item_commands,[])
        code=Form(gadgets=[d.gadget]).pml();self.assertIn('!this._Choice.rtext = !values',code)
        g=import_mac(code).form.gadgets[0];self.assertEqual(g.item_values,['',''])
        reopened=self.editor(g);self.assertEqual(reopened.mode.currentIndex(),1)
        reopened.accept();self.assertEqual(reopened.gadget.item_values,['',''])

    def test_blank_real_values_survive_unrelated_right_panel_edit(self):
        for kind in ('option','combo','list'):
            w=self.window(Form(gadgets=[Gadget(kind=kind,name='Choice',items=['A'],item_values=[''])]))
            try:
                w.fields['label'].setText('New title');w.update_gadget();self.assertEqual(w.form.gadgets[0].item_values,['']);w.form.validate()
            finally:w.dirty=False;w.close()

    def test_explicit_blank_b_column_is_real_data_for_combo_and_simple_list(self):
        for kind in ('combo','list'):
            d=self.editor(Gadget(kind=kind,name='Choice',items=['Old']))
            d.apply_excel_data([['Caption','Value'],['A',''],['B','']],first_row_header=True);d.accept()
            self.assertEqual(d.gadget.item_values,['',''])
            loaded=import_mac(Form(gadgets=[d.gadget]).pml()).form.gadgets[0];self.assertEqual(loaded.item_values,['',''])

    def test_one_column_excel_still_omits_real_values(self):
        d=self.editor(Gadget(kind='combo',name='Choice',items=['Old'],item_values=['']))
        d.apply_excel_data([['A']]);d.accept();self.assertEqual(d.gadget.item_values,[])

    def test_cell_validation_prevents_accepting_invalid_manual_edits(self):
        g=Gadget(kind='combo',name='Choice',items=['Original'],item_values=['1']);d=self.editor(g)
        d.table.setItem(0,0,QTableWidgetItem('Invalid\ncaption'));d.accept()
        self.assertEqual(d.result(),QDialog.Rejected);self.assertIn('A1',d.error.text());self.assertEqual(d.gadget,g)
        d.table.setItem(0,0,QTableWidgetItem('Corrected'));d.accept();self.assertEqual(d.result(),QDialog.Accepted)

    def test_invalid_initial_selection_stays_editable(self):
        d=self.editor(Gadget(kind='list',name='Rows',items=['A'],initial='2'));d.accept()
        self.assertEqual(d.result(),QDialog.Rejected);self.assertIn('初期選択',d.error.text());d.initial.setText('1');d.accept();self.assertEqual(d.result(),QDialog.Accepted)

    def test_column_address_after_z_is_excel_not_ambiguous_number(self):
        d=self.editor(Gadget(kind='list',name='Rows'))
        with self.assertRaisesRegex(ValueError,'AA1'):d.apply_excel_data([['A']*26+['bad\n']],True,True)

    def test_active_cell_editor_is_committed_by_ok(self):
        d=self.editor(Gadget(kind='combo',name='Choice',items=['Old']));d.show();self.app.processEvents()
        d.table.editItem(d.table.item(0,0));self.app.processEvents();editor=d.table.findChild(QLineEdit)
        self.assertIsNotNone(editor);editor.selectAll();QTest.keyClicks(editor,'Edited')
        buttons=d.findChild(QDialogButtonBox);QTest.mouseClick(buttons.button(QDialogButtonBox.Ok),Qt.LeftButton)
        self.assertEqual(d.gadget.items,['Edited'])

    def test_source_constructor_rejects_excel_replacement_before_table_mutates(self):
        for kind in ('combo','option','list'):
            g=Gadget(kind=kind,name='Choice',items=['Original'],item_values=['1'])
            d=self.editor(g,constructor_mode='SOURCE');before=d.table.item(0,0).text()
            with self.assertRaisesRegex(ValueError,'元コード優先'):d.apply_excel_data([['New','2']])
            self.assertEqual(d.table.item(0,0).text(),before);self.assertEqual(d.gadget,g)
            d.table.setItem(0,0,QTableWidgetItem('Manual'));d.accept();self.assertEqual(d.result(),QDialog.Rejected)

    def test_source_pair_commands_and_existing_table_helper_remain_editable(self):
        g=Gadget(kind='option',name='Choice',items=['A'],item_commands=['$P |A|'])
        d=self.editor(g,constructor_mode='SOURCE');d.apply_excel_data([['B','$P |B|']]);d.accept();self.assertEqual(d.result(),QDialog.Accepted)
        form=Form(constructor_mode='SOURCE',constructor_body="$P 'keep'",gadgets=[d.gadget]);self.assertIn("'B' '$P |B|'",form.pml())
        table=Gadget(kind='list',name='Rows',list_mode='TABLE',headings=['Name'],rows=[['A']],table_method='Populate')
        d=self.editor(table,constructor_mode='SOURCE');d.apply_excel_data([['Name'],['New']],True,True);d.accept();self.assertEqual(d.result(),QDialog.Accepted)
        form=Form(constructor_mode='SOURCE',constructor_body='!this.Populate()',gadgets=[d.gadget]);self.assertIn("= 'New'",form.pml())
        after=copy.deepcopy(table);after.table_method='Other'
        with self.assertRaisesRegex(ValueError,'元コード優先'):validate_item_change(table,after,'SOURCE')

    def test_source_change_from_pair_commands_to_real_values_is_rejected(self):
        d=self.editor(Gadget(kind='option',name='Choice',items=['A'],item_commands=['']),constructor_mode='SOURCE')
        d.mode.setCurrentIndex(1)
        with self.assertRaisesRegex(ValueError,'元コード優先'):d.apply_excel_data([['B','2']])

    def test_source_right_panel_edit_rolls_back_model_history_and_dirty_state(self):
        form=Form(constructor_mode='SOURCE',constructor_body="$P 'keep'",gadgets=[Gadget(kind='combo',name='Choice',items=['Original'])])
        w=self.window(form);w.dirty=False;w.future=[Form(title='future')];before=w.form.dumps();future=list(w.future)
        try:
            w.choices.blockSignals(True);w.choices.setPlainText('Ignored');w.choices.blockSignals(False);w.update_gadget()
            self.assertEqual(w.form.dumps(),before);self.assertEqual(w.history,[]);self.assertEqual(w.future,future);self.assertFalse(w.dirty)
            self.assertIn('元コード優先',w.statusBar().currentMessage());self.assertIn("$P 'keep'",w.form.pml())
        finally:w.dirty=False;w.close()

    def test_source_mini_editor_reports_block_and_retains_original(self):
        f=Form(constructor_mode='SOURCE',constructor_body="$P 'keep'",gadgets=[Gadget(kind='combo',name='Choice',items=['Original'])])
        outer=MiniProperties(None,f,0);self.dialogs.append(outer)
        def attempt(dialog):
            self.assertEqual(dialog.constructor_mode,'SOURCE');dialog.table.setItem(0,0,QTableWidgetItem('Ignored'));dialog.accept();self.assertIn('元コード優先',dialog.error.text());return dialog.result()
        with patch('e3d_designer.quick_editor.ItemsDialog.exec',attempt):outer.edit_items()
        self.assertEqual(outer.gadget.items,['Original']);self.assertEqual(f.gadgets[0].items,['Original'])

    def test_source_duplicate_with_required_initialization_is_transactional(self):
        form=Form(constructor_mode='SOURCE',constructor_body="$P 'keep'",gadgets=[
            Gadget(kind='frame',name='Group',width=50,height=20),Gadget(kind='combo',name='Choice',parent='Group',items=['Original'])])
        w=self.window(form);before=w.form.dumps();w.dirty=False
        try:
            w.duplicate();self.assertEqual(w.form.dumps(),before);self.assertEqual(w.history,[]);self.assertFalse(w.dirty)
            self.assertIn('Choice',w.statusBar().currentMessage())
            self.assertTrue(w.copy_gadget());w.paste_gadget();self.assertEqual(w.form.dumps(),before);self.assertEqual(w.history,[])
        finally:w.dirty=False;w.close()

    def test_source_allows_declaration_only_duplicate_and_cut_restore(self):
        g=Gadget(kind='option',name='Choice',items=['A'],item_commands=['$P |A|'])
        source=Form(constructor_mode='SOURCE',constructor_body="$P 'keep'",gadgets=[g]);copy_form,index=clone_subtree(source,source,0)
        self.assertEqual(copy_form.gadgets[index].item_commands,g.item_commands)
        combo=Gadget(kind='combo',name='Combo',items=['A']);source=Form(constructor_mode='SOURCE',constructor_body="$P 'keep'",gadgets=[combo])
        target=copy.deepcopy(source);target.gadgets=[];restored,index=clone_subtree(target,source,0,restore_names=True,same_project=True)
        self.assertEqual(restored.gadgets[index],combo);self.assertEqual(restored.constructor_body,source.constructor_body)

    def test_cut_restore_does_not_bypass_changed_source_initialization(self):
        source=Form(constructor_mode='SOURCE',constructor_body="!this.Combo.dtext = !data",gadgets=[Gadget(kind='combo',name='Combo',items=['A'])])
        target=copy.deepcopy(source);target.gadgets=[];target.constructor_body="$P 'changed'";before=target.dumps()
        with self.assertRaisesRegex(ValueError,'初期化'):clone_subtree(target,source,0,restore_names=True,same_project=True)
        self.assertEqual(target.dumps(),before)

    def test_source_constructor_required_controls_cannot_duplicate_without_initialization(self):
        gadgets=[Gadget(kind='list',name='Rows',list_mode='TABLE',headings=['A'],rows=[['1']]),
                 Gadget(kind='option',name='Choice',items=['A'],item_values=['1']),
                 Gadget(kind='button',name='Image',display_mode='PIXMAP',pixmap_path='sample.png',width=40,height=30),
                 Gadget(kind='slider',name='Level',callback='Changed',body="$P 'changed'")]
        for g in gadgets:
            source=Form(constructor_mode='SOURCE',constructor_body="$P 'keep'",gadgets=[g]);before=source.dumps()
            with self.assertRaisesRegex(ValueError,'元コード優先'):clone_subtree(source,source,0)
            self.assertEqual(source.dumps(),before)
